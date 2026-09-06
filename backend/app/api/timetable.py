import io
from typing import List, Optional, Dict, Any
from fastapi import APIRouter, Depends, HTTPException, status, Response
from sqlalchemy.orm import Session
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from reportlab.lib.pagesizes import letter, landscape
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors

from app.core.database import get_db
from app.models.all_models import (
    Timetable, TimetableEntry, GenerationRun, GenerationSolution, Subject, Teacher, Room, Section, Period, WorkingDay, AuditLog, Notification
)
from app.schemas.schemas import (
    TimetableResponse, TimetableEntryResponse, GenerateTimetableRequest,
    MoveClassRequest, SwapClassRequest, ValidationResult, GenerationRunResponse
)
from app.scheduler.input_parser import parse_schedule_problem
from app.scheduler.generator import generate_routine
from app.scheduler.validator import TimetableValidator
from app.scheduler.explainer import ScheduleExplainer
from app.scheduler.models import ScheduledLesson

router = APIRouter(prefix="/timetable", tags=["Timetable & Scheduling"])

@router.get("", response_model=List[TimetableResponse])
def get_timetables(db: Session = Depends(get_db)):
    timetables = db.query(Timetable).order_by(Timetable.id.desc()).all()
    res = []
    for tt in timetables:
        resp = TimetableResponse.from_orm(tt)
        resp.academic_year_name = tt.academic_year.name if tt.academic_year else None
        res.append(resp)
    return res

@router.get("/{id}", response_model=TimetableResponse)
def get_timetable(id: int, db: Session = Depends(get_db)):
    tt = db.query(Timetable).filter(Timetable.id == id).first()
    if not tt:
        raise HTTPException(status_code=404, detail="Timetable not found")

    resp = TimetableResponse.from_orm(tt)
    resp.academic_year_name = tt.academic_year.name if tt.academic_year else None
    
    entries_res = []
    for e in tt.entries:
        entry_resp = TimetableEntryResponse.from_orm(e)
        entry_resp.section_name = e.section.name if e.section else None
        entry_resp.program_name = e.section.semester.program.name if (e.section and e.section.semester and e.section.semester.program) else None
        entry_resp.subject_name = e.subject.name if e.subject else None
        entry_resp.subject_code = e.subject.code if e.subject else None
        entry_resp.subject_color = e.subject.color_code if e.subject else "#3B82F6"
        entry_resp.teacher_name = e.teacher.name if e.teacher else None
        entry_resp.teacher_designation = e.teacher.designation if e.teacher else None
        entry_resp.room_number = e.room.room_number if e.room else None
        entry_resp.room_type_name = e.room.room_type.name if (e.room and e.room.room_type) else None
        
        if e.period:
            entry_resp.period_name = e.period.name
            entry_resp.start_time = e.period.start_time
            entry_resp.end_time = e.period.end_time
            entry_resp.period_type = e.period.period_type
            if e.period.day:
                entry_resp.day_id = e.period.day.id
                entry_resp.day_name = e.period.day.name
                entry_resp.day_short_code = e.period.day.short_code

        entries_res.append(entry_resp)

    resp.entries = entries_res
    return resp

@router.post("/generate")
def generate_timetable_api(request: GenerateTimetableRequest, db: Session = Depends(get_db)):
    try:
        result = generate_routine(db=db, request=request)
        return result
    except Exception as e:
        raise HTTPException(status_code=400, detail=str(e))

from pydantic import BaseModel
from typing import List, Optional

class SimpleSubjectInput(BaseModel):
    name: str
    weekly_periods: int = 4
    teacher_name: str
    free_time_start: str = "08:00 AM"
    free_time_end: str = "04:00 PM"
    free_days: Optional[List[str]] = None

class SimplePeriodInput(BaseModel):
    name: str
    start_time: str
    end_time: str
    type: str = "Teaching"

class QuickWizardRequest(BaseModel):
    class_name: str = "BCA 1st Sem"
    routine_title: Optional[str] = None
    days: List[str] = ["Sunday", "Monday", "Tuesday", "Wednesday", "Thursday", "Friday"]
    periods: List[SimplePeriodInput]
    subjects: List[SimpleSubjectInput]
    day_period_counts: Optional[Dict[str, int]] = None
    day_periods: Optional[Dict[str, List[SimplePeriodInput]]] = None

@router.post("/quick-wizard")
def quick_wizard_api(req: QuickWizardRequest, db: Session = Depends(get_db)):
    from app.scheduler.simple_wizard import generate_simple_wizard_routine
    try:
        periods_dict = [p.dict() for p in req.periods]
        subjects_dict = [s.dict() for s in req.subjects]
        for s in subjects_dict:
            if not s.get("free_days"):
                s["free_days"] = req.days

        day_periods_dict = None
        if req.day_periods:
            day_periods_dict = {
                d: [p.dict() for p in plist] for d, plist in req.day_periods.items()
            }

        res = generate_simple_wizard_routine(
            db=db,
            class_name=req.class_name,
            days_list=req.days,
            periods_list=periods_dict,
            subjects_list=subjects_dict,
            routine_title=req.routine_title,
            day_periods=day_periods_dict,
            day_period_counts=req.day_period_counts
        )
        return res
    except Exception as e:
        raise HTTPException(status_code=400, detail=str(e))

class BCATeacherInput(BaseModel):
    name: str
    abbreviation: Optional[str] = None
    contact: Optional[str] = None
    phone: Optional[str] = None
    speciality: str = "Computer Science"
    free_time_start: str = "08:00 AM"
    free_time_end: str = "04:00 PM"
    free_days: Optional[List[str]] = None
    max_classes_per_day: int = 4

class BCASubjectInput(BaseModel):
    name: str
    code: Optional[str] = None
    weekly_periods: int = 4
    teacher_name: str
    course_type: Optional[str] = "TH"

class BCASemesterInput(BaseModel):
    semester_number: int
    semester_name: str = "BCA Semester 1"
    section_name: str = "BCA 1st Sem"
    room_name: str = "Room 101"
    subjects: List[BCASubjectInput]

class BCAMultiSemesterRequest(BaseModel):
    routine_title: Optional[str] = None
    days: List[str] = ["Sunday", "Monday", "Tuesday", "Wednesday", "Thursday", "Friday"]
    periods: List[SimplePeriodInput]
    day_period_counts: Optional[Dict[str, int]] = None
    teachers: List[BCATeacherInput]
    running_semesters: List[BCASemesterInput]

@router.post("/bca-routine")
def generate_bca_routine_api(req: BCAMultiSemesterRequest, db: Session = Depends(get_db)):
    from app.scheduler.bca_multi_semester import generate_bca_multi_semester_routine
    try:
        teachers_dict = [t.model_dump() if hasattr(t, 'model_dump') else t.dict() for t in req.teachers]
        for t in teachers_dict:
            if not t.get("free_days"):
                t["free_days"] = req.days

        semesters_dict = [
            {
                "semester_number": s.semester_number,
                "semester_name": s.semester_name,
                "section_name": s.section_name,
                "room_name": s.room_name,
                "subjects": [sub.model_dump() if hasattr(sub, 'model_dump') else sub.dict() for sub in s.subjects]
            }
            for s in req.running_semesters
        ]

        periods_dict = [p.model_dump() if hasattr(p, 'model_dump') else p.dict() for p in req.periods]

        result = generate_bca_multi_semester_routine(
            db=db,
            running_semesters=semesters_dict,
            teachers_list=teachers_dict,
            days_list=req.days,
            periods_list=periods_dict,
            day_period_counts=req.day_period_counts,
            routine_title=req.routine_title
        )
        return result
    except Exception as e:
        raise HTTPException(status_code=400, detail=str(e))


@router.post("/validate", response_model=ValidationResult)
def validate_timetable_api(timetable_id: int, db: Session = Depends(get_db)):
    tt = db.query(Timetable).filter(Timetable.id == timetable_id).first()
    if not tt:
        raise HTTPException(status_code=404, detail="Timetable not found")

    problem = parse_schedule_problem(db=db)
    validator = TimetableValidator(problem)

    lessons = [
        ScheduledLesson(
            section_id=e.section_id,
            subject_id=e.subject_id,
            teacher_id=e.teacher_id,
            room_id=e.room_id,
            period_id=e.period_id,
            day_id=e.period.day_id if e.period else 1,
            is_locked=e.is_locked
        )
        for e in tt.entries
    ]

    val_res = validator.validate_full(lessons)
    # Update conflict count in timetable
    tt.conflict_count = val_res.total_conflicts
    db.commit()
    return val_res

@router.post("/move")
def move_class_slot(req: MoveClassRequest, db: Session = Depends(get_db)):
    entry = db.query(TimetableEntry).filter(
        TimetableEntry.id == req.entry_id,
        TimetableEntry.timetable_id == req.timetable_id
    ).first()
    if not entry:
        raise HTTPException(status_code=404, detail="Timetable entry not found")

    target_period = db.query(Period).filter(Period.id == req.target_period_id).first()
    if not target_period:
        raise HTTPException(status_code=404, detail="Target period not found")

    target_room_id = req.target_room_id or entry.room_id

    # Build lessons list to simulate
    problem = parse_schedule_problem(db=db)
    validator = TimetableValidator(problem)

    tt = db.query(Timetable).filter(Timetable.id == req.timetable_id).first()
    entries = list(tt.entries)
    entry_idx = next((i for i, e in enumerate(entries) if e.id == req.entry_id), -1)

    lessons = [
        ScheduledLesson(
            section_id=e.section_id,
            subject_id=e.subject_id,
            teacher_id=e.teacher_id,
            room_id=e.room_id,
            period_id=e.period_id,
            day_id=e.period.day_id if e.period else 1
        )
        for e in entries
    ]

    is_allowed, msg, val_res = validator.validate_move(
        lessons, entry_idx, req.target_period_id, target_room_id
    )

    if not is_allowed:
        raise HTTPException(status_code=400, detail=f"Move not allowed: {msg}")

    # Apply move
    old_p_name = entry.period.name if entry.period else ""
    entry.period_id = req.target_period_id
    entry.room_id = target_room_id
    tt.updated_at = db.query(Timetable).filter(Timetable.id == req.timetable_id).first().updated_at
    
    # Audit log
    audit = AuditLog(
        action="MOVE_CLASS",
        entity_type="TimetableEntry",
        entity_id=entry.id,
        details={
            "subject": entry.subject.name if entry.subject else "",
            "old_period": old_p_name,
            "new_period": target_period.name,
            "new_day": target_period.day.name if target_period.day else ""
        }
    )
    db.add(audit)
    db.commit()

    return {"status": "success", "message": "Class moved successfully", "validation": val_res}

@router.post("/swap")
def swap_class_slots(req: SwapClassRequest, db: Session = Depends(get_db)):
    entry_a = db.query(TimetableEntry).filter(
        TimetableEntry.id == req.entry_a_id,
        TimetableEntry.timetable_id == req.timetable_id
    ).first()
    entry_b = db.query(TimetableEntry).filter(
        TimetableEntry.id == req.entry_b_id,
        TimetableEntry.timetable_id == req.timetable_id
    ).first()

    if not entry_a or not entry_b:
        raise HTTPException(status_code=404, detail="One or both timetable entries not found")

    problem = parse_schedule_problem(db=db)
    validator = TimetableValidator(problem)

    tt = db.query(Timetable).filter(Timetable.id == req.timetable_id).first()
    entries = list(tt.entries)
    idx_a = next((i for i, e in enumerate(entries) if e.id == req.entry_a_id), -1)
    idx_b = next((i for i, e in enumerate(entries) if e.id == req.entry_b_id), -1)

    lessons = [
        ScheduledLesson(
            section_id=e.section_id,
            subject_id=e.subject_id,
            teacher_id=e.teacher_id,
            room_id=e.room_id,
            period_id=e.period_id,
            day_id=e.period.day_id if e.period else 1
        )
        for e in entries
    ]

    is_allowed, msg, val_res = validator.validate_swap(lessons, idx_a, idx_b)

    if not is_allowed:
        raise HTTPException(status_code=400, detail=f"Swap not allowed: {msg}")

    # Apply swap
    p_a = entry_a.period_id
    r_a = entry_a.room_id
    entry_a.period_id = entry_b.period_id
    entry_a.room_id = entry_b.room_id
    entry_b.period_id = p_a
    entry_b.room_id = r_a

    db.commit()
    return {"status": "success", "message": "Classes swapped successfully"}

@router.get("/{id}/explain/{entry_id}")
def explain_entry(id: int, entry_id: int, db: Session = Depends(get_db)):
    entry = db.query(TimetableEntry).filter(
        TimetableEntry.id == entry_id,
        TimetableEntry.timetable_id == id
    ).first()
    if not entry:
        raise HTTPException(status_code=404, detail="Timetable entry not found")

    problem = parse_schedule_problem(db=db)
    explainer = ScheduleExplainer(problem)

    lesson = ScheduledLesson(
        section_id=entry.section_id,
        subject_id=entry.subject_id,
        teacher_id=entry.teacher_id,
        room_id=entry.room_id,
        period_id=entry.period_id,
        day_id=entry.period.day_id if entry.period else 1
    )

    explanation = explainer.explain_lesson(lesson)
    return explanation

@router.post("/{id}/publish")
def publish_timetable(id: int, db: Session = Depends(get_db)):
    tt = db.query(Timetable).filter(Timetable.id == id).first()
    if not tt:
        raise HTTPException(status_code=404, detail="Timetable not found")

    tt.is_published = True
    notif = Notification(
        title="Timetable Published",
        message=f"Timetable '{tt.name}' has been officially published and is now active.",
        notification_type="ROUTINE_UPDATE"
    )
    db.add(notif)
    db.commit()
    return {"status": "success", "message": "Timetable published successfully"}

@router.post("/{id}/apply-solution/{solution_id}")
def apply_candidate_solution(id: int, solution_id: int, db: Session = Depends(get_db)):
    tt = db.query(Timetable).filter(Timetable.id == id).first()
    if not tt:
        raise HTTPException(status_code=404, detail="Timetable not found")

    sol = db.query(GenerationSolution).filter(GenerationSolution.id == solution_id).first()
    if not sol:
        raise HTTPException(status_code=404, detail="Candidate solution not found")

    # Clear current entries
    db.query(TimetableEntry).filter(TimetableEntry.timetable_id == id).delete()

    problem = parse_schedule_problem(db=db)
    explainer = ScheduleExplainer(problem)

    for item in sol.entries_data:
        lesson = ScheduledLesson(
            section_id=item["section_id"],
            subject_id=item["subject_id"],
            teacher_id=item["teacher_id"],
            room_id=item["room_id"],
            period_id=item["period_id"],
            day_id=item["day_id"]
        )
        exp = explainer.explain_lesson(lesson)
        entry = TimetableEntry(
            timetable_id=tt.id,
            section_id=lesson.section_id,
            subject_id=lesson.subject_id,
            teacher_id=lesson.teacher_id,
            room_id=lesson.room_id,
            period_id=lesson.period_id,
            explanation=exp.get("summary", "")
        )
        db.add(entry)

    tt.score = sol.score
    tt.version += 1
    db.commit()
    return {"status": "success", "message": f"Applied solution #{sol.solution_index} with score {sol.score}%"}

# --- Export to Excel --- #
@router.get("/{id}/export/excel")
def export_timetable_excel(id: int, section_id: Optional[int] = None, db: Session = Depends(get_db)):
    tt = db.query(Timetable).filter(Timetable.id == id).first()
    if not tt:
        raise HTTPException(status_code=404, detail="Timetable not found")

    campus = db.query(Campus).first()
    campus_name = campus.name if campus else "Ratna Rajyalaxmi Campus"
    campus_addr = campus.address if (campus and campus.address) else "Pradarshanimarga, Kathmandu Nepal"

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Campus Routine"

    title_font = Font(name="Arial", size=13, bold=True)
    sub_font = Font(name="Arial", size=10, bold=True)
    header_font = Font(name="Arial", size=10, bold=True)
    header_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    interval_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    thin_border_side = Side(border_style="thin", color="000000")
    thin_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
    center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

    days = db.query(WorkingDay).filter(WorkingDay.is_active == True).order_by(WorkingDay.order_index).all()
    sections = db.query(Section).all()
    if section_id:
        sections = [s for s in sections if s.id == section_id]

    all_entries = db.query(TimetableEntry).filter(TimetableEntry.timetable_id == id).all()

    current_row = 1
    for sec in sections:
        # College Header
        prog_name = sec.semester.program.name if (sec.semester and sec.semester.program) else "Bachelors in Computer Applications (BCA)"
        sem_name = sec.semester.name if sec.semester else ""
        room_name = ""
        sec_entries = [e for e in all_entries if e.section_id == sec.id]
        if sec_entries and sec_entries[0].room:
            room_name = sec_entries[0].room.room_number

        ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=7)
        ws.cell(row=current_row, column=1, value=campus_name).font = title_font
        ws.cell(row=current_row, column=1).alignment = center_align
        current_row += 1

        ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=7)
        ws.cell(row=current_row, column=1, value=campus_addr).font = sub_font
        ws.cell(row=current_row, column=1).alignment = center_align
        current_row += 1

        ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=7)
        ws.cell(row=current_row, column=1, value=f"{prog_name} {sem_name} - Sec {sec.name} Room No- {room_name}").font = sub_font
        ws.cell(row=current_row, column=1).alignment = center_align
        current_row += 1

        ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=7)
        ws.cell(row=current_row, column=1, value="Daily Class Routine").font = sub_font
        ws.cell(row=current_row, column=1).alignment = center_align
        current_row += 1

        # Table Header
        ws.cell(row=current_row, column=1, value="Day/Time").font = header_font
        ws.cell(row=current_row, column=1).fill = header_fill
        ws.cell(row=current_row, column=1).border = thin_border
        ws.cell(row=current_row, column=1).alignment = center_align

        # Get representative periods from first day
        rep_periods = days[0].periods if days else []
        for p_idx, p in enumerate(rep_periods[:6]):
            col = p_idx + 2
            cell = ws.cell(row=current_row, column=col, value=f"{p.start_time}-{p.end_time}")
            cell.font = header_font
            cell.fill = header_fill
            cell.border = thin_border
            cell.alignment = center_align

        current_row += 1
        table_start_row = current_row

        teachers_in_sec = {}

        for day in days:
            ws.cell(row=current_row, column=1, value=day.name).font = header_font
            ws.cell(row=current_row, column=1).fill = header_fill
            ws.cell(row=current_row, column=1).border = thin_border
            ws.cell(row=current_row, column=1).alignment = center_align

            for p_idx, p in enumerate(day.periods[:6]):
                col = p_idx + 2
                cell = ws.cell(row=current_row, column=col)
                cell.border = thin_border
                cell.alignment = center_align

                if p.period_type != "Teaching":
                    cell.value = "Interval"
                    cell.fill = interval_fill
                else:
                    match = next((e for e in all_entries if e.section_id == sec.id and e.period_id == p.id), None)
                    if match:
                        t_name = match.teacher.name if match.teacher else ""
                        words = [w for w in re.split(r'[\s._-]+', t_name) if w and w.lower() not in ["prof", "dr", "er", "mr", "mrs", "ms"]]
                        t_abbrev = "".join(w[0] for w in (words if words else [t_name])).upper() if t_name else "TCH"
                        if match.teacher:
                            teachers_in_sec[t_abbrev] = {
                                "abbrev": t_abbrev,
                                "name": match.teacher.name,
                                "contact": match.teacher.phone or "9841000000"
                            }
                        cell.value = f"{match.subject.name}\n[TH] [{t_abbrev}]"
                    else:
                        cell.value = ""

            current_row += 1

        current_row += 1

        # Teacher Abbreviation Table & Legend
        ws.cell(row=current_row, column=1, value="Abbrev").font = header_font
        ws.cell(row=current_row, column=1).border = thin_border
        ws.cell(row=current_row, column=2, value="Name").font = header_font
        ws.cell(row=current_row, column=2).border = thin_border
        ws.cell(row=current_row, column=3, value="Contact").font = header_font
        ws.cell(row=current_row, column=3).border = thin_border

        ws.cell(row=current_row, column=5, value="TH=Theory,").font = Font(name="Arial", size=9)
        ws.cell(row=current_row, column=6, value="TU=Tutorial").font = Font(name="Arial", size=9)
        current_row += 1

        for t_info in teachers_in_sec.values():
            ws.cell(row=current_row, column=1, value=t_info["abbrev"]).border = thin_border
            ws.cell(row=current_row, column=2, value=t_info["name"]).border = thin_border
            ws.cell(row=current_row, column=3, value=t_info["contact"]).border = thin_border
            if t_info == list(teachers_in_sec.values())[0]:
                ws.cell(row=current_row, column=5, value="PR=Practical").font = Font(name="Arial", size=9)
            current_row += 1

        current_row += 3

    for col in ws.columns:
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = 18

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    headers = {"Content-Disposition": f"attachment; filename=routine_{id}.xlsx"}
    return Response(content=output.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers=headers)

# --- Export to PDF --- #
@router.get("/{id}/export/pdf")
def export_timetable_pdf(id: int, section_id: Optional[int] = None, db: Session = Depends(get_db)):
    tt = db.query(Timetable).filter(Timetable.id == id).first()
    if not tt:
        raise HTTPException(status_code=404, detail="Timetable not found")

    campus = db.query(Campus).first()
    campus_name = campus.name if campus else "Ratna Rajyalaxmi Campus"
    campus_addr = campus.address if (campus and campus.address) else "Pradarshanimarga, Kathmandu Nepal"

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=landscape(letter), rightMargin=24, leftMargin=24, topMargin=24, bottomMargin=24)
    story = []
    styles = getSampleStyleSheet()

    title_style = ParagraphStyle("CampusTitle", parent=styles["Normal"], fontName="Helvetica-Bold", fontSize=14, leading=16, alignment=1)
    sub_style = ParagraphStyle("CampusSub", parent=styles["Normal"], fontName="Helvetica", fontSize=10, leading=13, alignment=1)
    routine_title_style = ParagraphStyle("RoutineTitle", parent=styles["Normal"], fontName="Helvetica-Bold", fontSize=11, leading=14, alignment=1)

    days = db.query(WorkingDay).filter(WorkingDay.is_active == True).order_by(WorkingDay.order_index).all()
    sections = db.query(Section).all()
    if section_id:
        sections = [s for s in sections if s.id == section_id]

    all_entries = db.query(TimetableEntry).filter(TimetableEntry.timetable_id == id).all()

    for sec in sections:
        prog_name = sec.semester.program.name if (sec.semester and sec.semester.program) else "Bachelors in Computer Applications (BCA)"
        sem_name = sec.semester.name if sec.semester else ""
        room_name = ""
        sec_entries = [e for e in all_entries if e.section_id == sec.id]
        if sec_entries and sec_entries[0].room:
            room_name = sec_entries[0].room.room_number

        story.append(Paragraph(f"<b>{campus_name}</b>", title_style))
        story.append(Paragraph(f"{campus_addr}", sub_style))
        story.append(Paragraph(f"{prog_name} {sem_name} - Sec {sec.name} Room No- {room_name}", sub_style))
        story.append(Paragraph("<b>Daily Class Routine</b>", routine_title_style))
        story.append(Spacer(1, 8))

        # Build table data
        rep_periods = days[0].periods if days else []
        header_row = ["Day/Time"] + [f"{p.start_time}-{p.end_time}" for p in rep_periods[:6]]
        table_data = [header_row]

        teachers_in_sec = {}

        for day in days:
            row = [day.name]
            for p_idx, p in enumerate(day.periods[:6]):
                if p.period_type != "Teaching":
                    row.append("Interval")
                else:
                    match = next((e for e in all_entries if e.section_id == sec.id and e.period_id == p.id), None)
                    if match:
                        t_name = match.teacher.name if match.teacher else ""
                        words = [w for w in re.split(r'[\s._-]+', t_name) if w and w.lower() not in ["prof", "dr", "er", "mr", "mrs", "ms"]]
                        t_abbrev = "".join(w[0] for w in (words if words else [t_name])).upper() if t_name else "TCH"
                        if match.teacher:
                            teachers_in_sec[t_abbrev] = {
                                "abbrev": t_abbrev,
                                "name": match.teacher.name,
                                "contact": match.teacher.phone or "9841000000"
                            }
                        row.append(f"{match.subject.name}\n[TH] [{t_abbrev}]")
                    else:
                        row.append("")
            table_data.append(row)

        pdf_table = Table(table_data, colWidths=[70] + [95] * (len(header_row) - 1))
        pdf_table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#F1F5F9")),
            ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#F1F5F9")),
            ("TEXTCOLOR", (0, 0), (-1, -1), colors.HexColor("#0F172A")),
            ("ALIGN", (0, 0), (-1, -1), "CENTER"),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 8),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#000000")),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ("TOPPADDING", (0, 0), (-1, -1), 4),
        ]))
        story.append(pdf_table)
        story.append(Spacer(1, 10))

        # Teacher Abbrev & Legend Table
        t_rows = [["Abbrev", "Name", "Contact", "", "Legend"]]
        t_list = list(teachers_in_sec.values())
        for idx, t_info in enumerate(t_list):
            legend_text = ""
            if idx == 0:
                legend_text = "TH=Theory, TU=Tutorial"
            elif idx == 1:
                legend_text = "PR=Practical"
            t_rows.append([t_info["abbrev"], t_info["name"], t_info["contact"], "", legend_text])

        if len(t_list) == 0:
            t_rows.append(["-", "No teacher assigned", "-", "", "TH=Theory, PR=Practical"])

        t_table = Table(t_rows, colWidths=[60, 160, 90, 40, 160])
        t_table.setStyle(TableStyle([
            ("FONTNAME", (0, 0), (2, 0), "Helvetica-Bold"),
            ("GRID", (0, 0), (2, -1), 0.5, colors.HexColor("#000000")),
            ("BACKGROUND", (0, 0), (2, 0), colors.HexColor("#F1F5F9")),
            ("FONTSIZE", (0, 0), (-1, -1), 8),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("ALIGN", (0, 0), (0, -1), "CENTER"),
        ]))
        story.append(t_table)
        story.append(Spacer(1, 20))

    doc.build(story)
    buffer.seek(0)
    headers = {"Content-Disposition": f"attachment; filename=routine_{id}.pdf"}
    return Response(content=buffer.getvalue(), media_type="application/pdf", headers=headers)
